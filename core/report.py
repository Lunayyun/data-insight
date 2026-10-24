"""Downloadable report generation: Markdown and PPTX (16:9, Lieflat-styled)."""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime
from typing import Iterable

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from core.charts import STYLE_PRESETS
from core.profiler import DataProfile


# ---- Layout tokens (16:9, 13.333 x 7.5 in) ----
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
TITLE_H = 1.0
FOOTER_H = 0.4
CONTENT_TOP = MARGIN + TITLE_H
CONTENT_H = SLIDE_H - CONTENT_TOP - FOOTER_H - 0.3

FONT_FAMILY = "Inter"


def _theme(preset: str) -> dict:
    return STYLE_PRESETS.get(preset, STYLE_PRESETS["Porcelain"])


def _rgb(hex_color: str) -> RGBColor:
    """Convert #hex or rgba(r,g,b,a) to RGBColor (alpha dropped for pptx)."""
    c = str(hex_color).strip()
    if c.startswith("rgba") or c.startswith("rgb"):
        inner = c[c.index("(") + 1 : c.rindex(")")]
        parts = [float(p.strip()) for p in inner.split(",")]
        return RGBColor(int(parts[0]), int(parts[1]), int(parts[2]))
    h = c.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_bg(slide, color: str):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=None, align=PP_ALIGN.LEFT, font=FONT_FAMILY):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    if color:
        run.font.color.rgb = _rgb(color)
    return tb


def _add_accent_bar(slide, theme):
    """A thin primary-color accent bar under the title (Lieflat hairline)."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(CONTENT_TOP - 0.18), Inches(1.2), Inches(0.05)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(theme["primary"])
    shape.line.fill.background()
    return shape


def _add_footer(slide, theme):
    _add_textbox(slide, MARGIN, SLIDE_H - FOOTER_H, 8, 0.3, "来源：DataInsight",
                 size=9, color=theme["muted"])


@dataclass
class SlideChart:
    image_bytes: bytes
    title: str
    caption: str = ""


@dataclass
class SlideTable:
    df: pd.DataFrame
    title: str
    image_bytes: bytes | None = None


def build_markdown_report(df: pd.DataFrame, profile: DataProfile, title: str) -> str:
    missing = "无" if profile.missing_summary.empty else profile.missing_summary.to_markdown(index=False)
    return f"""# {title}

生成时间：{datetime.now():%Y-%m-%d %H:%M}

## 数据摘要

- 数据行数：{profile.rows:,}
- 字段数量：{profile.columns:,}
- 缺失单元格：{profile.missing_cells:,}
- 重复行：{profile.duplicate_rows:,}

## 字段概况

{profile.column_summary.to_markdown(index=False)}

## 缺失值明细

{missing}

## 当前数据样例

{df.head(20).to_markdown(index=False)}
"""


def build_pptx_report(
    title: str,
    profile: DataProfile,
    charts: Iterable[SlideChart],
    tables: Iterable[SlideTable] | None = None,
    insights: list[str] | None = None,
    preset: str = "Porcelain",
) -> bytes:
    """Build a 16:9 PPTX and return the file bytes."""
    theme = _theme(preset)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    # ---- Cover slide ----
    cover = prs.slides.add_slide(blank)
    _set_bg(cover, theme["background"])
    accent = cover.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(SLIDE_W), Inches(0.25))
    accent.fill.solid(); accent.fill.fore_color.rgb = _rgb(theme["primary"]); accent.line.fill.background()
    _add_textbox(cover, MARGIN, 2.4, 11, 1.2, title, size=40, bold=True, color=theme["ink"])
    _add_textbox(cover, MARGIN, 3.7, 11, 0.6, "DataInsight · 数据分析报告", size=18, color=theme["muted"])
    _add_textbox(cover, MARGIN, 6.6, 11, 0.4, f"{datetime.now():%Y-%m-%d}", size=12, color=theme["muted"])

    # ---- Summary slide ----
    summary = prs.slides.add_slide(blank)
    _set_bg(summary, theme["background"])
    _add_textbox(summary, MARGIN, MARGIN, 11, 0.8, "数据摘要", size=28, bold=True, color=theme["ink"])
    _add_accent_bar(summary, theme)
    metrics = [
        ("有效记录", f"{profile.rows:,}"),
        ("字段数", f"{profile.columns:,}"),
        ("缺失单元格", f"{profile.missing_cells:,}"),
        ("重复行", f"{profile.duplicate_rows:,}"),
    ]
    card_w, card_h, gap = 2.8, 1.6, 0.25
    total_w = len(metrics) * card_w + (len(metrics) - 1) * gap
    start_x = (SLIDE_W - total_w) / 2
    for i, (label, value) in enumerate(metrics):
        x = start_x + i * (card_w + gap)
        card = summary.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(CONTENT_TOP + 0.4), Inches(card_w), Inches(card_h))
        card.fill.solid(); card.fill.fore_color.rgb = _rgb(theme["background"])
        card.line.color.rgb = _rgb(theme["grid"]); card.line.width = Pt(1)
        _add_textbox(summary, x + 0.2, CONTENT_TOP + 0.55, card_w - 0.4, 0.4, label, size=12, color=theme["muted"], align=PP_ALIGN.CENTER)
        _add_textbox(summary, x + 0.2, CONTENT_TOP + 0.95, card_w - 0.4, 0.8, value, size=32, bold=True, color=theme["primary"], align=PP_ALIGN.CENTER)
    if insights:
        _add_textbox(summary, MARGIN, CONTENT_TOP + 2.4, 11, 0.5, "AI 关键发现", size=16, bold=True, color=theme["ink"])
        insight_text = "\n".join(f"• {ins}" for ins in insights[:4])
        _add_textbox(summary, MARGIN, CONTENT_TOP + 2.95, 11, 2.5, insight_text, size=13, color=theme["ink"])
    _add_footer(summary, theme)

    # ---- Chart slides ----
    for chart in charts:
        slide = prs.slides.add_slide(blank)
        _set_bg(slide, theme["background"])
        _add_textbox(slide, MARGIN, MARGIN, 11, 0.8, chart.title, size=24, bold=True, color=theme["ink"])
        _add_accent_bar(slide, theme)
        img_stream = io.BytesIO(chart.image_bytes)
        slide.shapes.add_picture(img_stream, Inches(MARGIN), Inches(CONTENT_TOP), width=Inches(SLIDE_W - 2 * MARGIN))
        if chart.caption:
            _add_textbox(slide, MARGIN, SLIDE_H - FOOTER_H - 0.3, 11, 0.3, chart.caption, size=10, color=theme["muted"])
        _add_footer(slide, theme)

    # ---- Table slides ----
    for table in tables or []:
        slide = prs.slides.add_slide(blank)
        _set_bg(slide, theme["background"])
        _add_textbox(slide, MARGIN, MARGIN, 11, 0.8, table.title, size=24, bold=True, color=theme["ink"])
        _add_accent_bar(slide, theme)
        if table.image_bytes:
            slide.shapes.add_picture(io.BytesIO(table.image_bytes), Inches(MARGIN), Inches(CONTENT_TOP), width=Inches(SLIDE_W - 2 * MARGIN))
        _add_footer(slide, theme)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
